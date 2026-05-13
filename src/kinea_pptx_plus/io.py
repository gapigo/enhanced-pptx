from __future__ import annotations

from pptx import Presentation
from pptx.chart.chart import Chart
from pptx.shapes.graphfrm import GraphicFrame

import lxml.etree as etree


def extract_chart_xml(pptx_path: str, chart_index: int = 0) -> bytes:
    """Devolve o XML cru do chart no índice especificado.

    Lê o pptx via python-pptx-ng, encontra todos os GraphicFrames com has_chart,
    pega o de índice <chart_index>, retorna o XML serializado do chartSpace.

    Example:
        >>> xml = extract_chart_xml("fixture.pptx", 0)
        >>> b"<c:chartSpace" in xml
        True
    """
    prs = Presentation(pptx_path)

    charts: list[Chart] = []
    for slide in prs.slides:
        for shape in slide.shapes:
            if isinstance(shape, GraphicFrame) and shape.has_chart:
                charts.append(shape.chart)

    if chart_index >= len(charts):
        raise IndexError(
            f"chart_index={chart_index} out of range; "
            f"{len(charts)} chart(s) found"
        )

    chart: Chart = charts[chart_index]
    # _chartSpace is the internal lxml element for the chart XML
    chart_space = chart._chartSpace  # type: ignore[attr-defined]
    return etree.tostring(chart_space, xml_declaration=True, encoding="UTF-8")
