"""Test overlays.label_box — coloured rectangle labels."""

from __future__ import annotations

import lxml.etree as etree
from pptx import Presentation

from kinea_pptx_plus.overlays.label_box import add_label_box, _build_label_sp


def test_build_label_sp_creates_valid_xml():
    """The generated p:sp XML should parse cleanly."""
    sp = _build_label_sp(
        x=5000000, y=3000000,
        cx=1200000, cy=400000,
        fill_hex="1F3864", text="42.5%", font_size_pt=9,
    )
    xml = etree.tostring(sp, pretty_print=True)
    assert b"p:sp" in xml or b"p:sp" in xml or b"{http://schemas.openxmlformats.org/presentationml/2006/main}sp" in xml
    # Verify fill colour
    assert b"1F3864" in xml
    # Verify text
    assert b"42.5%" in xml


def test_build_label_sp_white_text_on_dark_bg():
    """Dark blue bg should get white text."""
    sp = _build_label_sp(
        x=0, y=0, cx=1000000, cy=300000,
        fill_hex="1F3864", text="10%", font_size_pt=9,
    )
    xml = etree.tostring(sp)
    # Should have white text (FFFFFF)
    assert b"FFFFFF" in xml


def test_build_label_sp_black_text_on_light_bg():
    """Light yellow bg should get black text."""
    sp = _build_label_sp(
        x=0, y=0, cx=1000000, cy=300000,
        fill_hex="F0E68C", text="10%", font_size_pt=9,
    )
    xml = etree.tostring(sp)
    # Should have black text (000000)
    assert b"000000" in xml


def test_build_label_sp_positions_correctly():
    """Check EMU coordinates are stored in a:off."""
    sp = _build_label_sp(
        x=5000000, y=3000000,
        cx=1200000, cy=400000,
        fill_hex="C0392B", text="Test", font_size_pt=9,
    )
    assert b"5000000" in etree.tostring(sp)
    assert b"3000000" in etree.tostring(sp)


def test_add_label_box_to_slide():
    """Add a label to a real slide and verify PPTX saves cleanly."""
    prs = Presentation()
    prs.slide_width = 9144000
    prs.slide_height = 6858000
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    add_label_box(
        slide,
        x_emu=5000000, y_emu=3000000,
        cx_emu=1200000, cy_emu=400000,
        fill_hex="1F3864", text="42.5%",
    )

    # Save and verify
    import tempfile, os
    with tempfile.NamedTemporaryFile(suffix=".pptx", delete=False) as tmp:
        out_path = tmp.name
    try:
        prs.save(out_path)
        # Reload and check
        prs2 = Presentation(out_path)
        slide2 = prs2.slides[0]
        assert slide2.shapes is not None
        assert len(slide2.shapes) >= 1
    finally:
        if os.path.exists(out_path):
            os.unlink(out_path)


def test_label_box_kinea_colors():
    """Test with Kinea brand colours."""
    for fill_hex, expected_text_color in [
        ("1F3864", "FFFFFF"),  # dark blue → white
        ("C0392B", "FFFFFF"),  # red → white
        ("F0E68C", "000000"),  # light → black
    ]:
        sp = _build_label_sp(
            0, 0, 1000000, 300000,
            fill_hex=fill_hex, text="99.9%", font_size_pt=9,
        )
        xml = etree.tostring(sp)
        assert expected_text_color.encode() in xml, (
            f"Expected {expected_text_color} for fill {fill_hex}"
        )
