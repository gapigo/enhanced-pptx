"""Test overlays.column_annotation — text box with leader line."""

from __future__ import annotations

import os
import tempfile

import lxml.etree as etree
from pptx import Presentation

from kinea_pptx_plus.overlays.column_annotation import (
    add_column_annotation,
    _build_annotation_sp,
    _build_leader_cxn_sp,
)


def test_build_annotation_sp_creates_valid_xml():
    sp = _build_annotation_sp(
        x=5000000, y=3000000, cx=2000000, cy=500000,
        text="Key insight", font_size_pt=10, fill_hex="F2F2F2",
    )
    xml = etree.tostring(sp)
    assert b"Key insight" in xml
    assert b"F2F2F2" in xml


def test_build_leader_cxn_sp():
    cxn = _build_leader_cxn_sp(5000000, 3000000, 6000000, 4000000)
    xml = etree.tostring(cxn)
    assert b"cxnSp" in xml or b"connector" in xml.lower() or b"cxn" in xml
    assert b"line" in xml


def test_add_column_annotation_to_slide():
    prs = Presentation()
    prs.slide_width = 9144000
    prs.slide_height = 6858000
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    add_column_annotation(
        slide,
        x_emu=5000000, y_emu=3000000,
        cx_emu=2000000, cy_emu=500000,
        text="Key insight",
        target_x=6000000, target_y=4000000,
    )

    with tempfile.NamedTemporaryFile(suffix=".pptx", delete=False) as tmp:
        out_path = tmp.name
    try:
        prs.save(out_path)
        prs2 = Presentation(out_path)
        assert len(prs2.slides) == 1
    finally:
        if os.path.exists(out_path):
            os.unlink(out_path)
