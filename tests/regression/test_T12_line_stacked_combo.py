"""Test line_stacked_combo — stacked bar + line with annotation."""

from __future__ import annotations

import os
import tempfile

import lxml.etree as etree
from pptx import Presentation

from kinea_pptx_plus.charts.line_stacked_combo import build_line_stacked_combo
from kinea_pptx_plus.io import extract_chart_xml


def _run_combo(path: str, **kw) -> str:
    with tempfile.NamedTemporaryFile(suffix=".pptx", delete=False) as tmp:
        out_path = tmp.name
    try:
        result = build_line_stacked_combo(path, output_path=out_path, **kw)
        assert result == out_path
        return out_path
    except Exception:
        if os.path.exists(out_path):
            os.unlink(out_path)
        raise


def test_combo_roundtrip_simple():
    """Simple roundtrip with the simplest fixture."""
    out = _run_combo("tests/fixtures/simple_line_chart.pptx")
    try:
        prs = Presentation(out)
        assert len(prs.slides) == 1
    finally:
        if os.path.exists(out):
            os.unlink(out)


def test_combo_schema_gate():
    """schema gate: chart XML must parse cleanly."""
    out = _run_combo("tests/fixtures/simple_line_chart.pptx")
    try:
        xml = extract_chart_xml(out)
        root = etree.fromstring(xml)
        assert root.tag.endswith("chartSpace")
    finally:
        if os.path.exists(out):
            os.unlink(out)


def test_combo_with_annotation():
    """Build combo chart with a callout annotation."""
    out = _run_combo(
        "tests/fixtures/simple_line_chart.pptx",
        annotation_text="Peak value",
        annotation_target_cat=3,
    )
    try:
        prs = Presentation(out)
        slide = prs.slides[0]
        assert len(slide.shapes) >= 1
    finally:
        if os.path.exists(out):
            os.unlink(out)
